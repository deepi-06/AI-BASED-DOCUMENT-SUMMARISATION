"""
document_reader.py
------------------
Extracts plain text from PDF, DOCX and TXT files.

Every problem (empty file, corrupt file, scanned PDF, ...) is raised as a
DocumentError with a friendly message that can be shown directly to the user.
"""

import os

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentError(Exception):
    """Raised when a document cannot be read or contains no usable text."""


def get_extension(filename: str) -> str:
    """Return the lower-case extension of a filename, e.g. '.pdf'."""
    return os.path.splitext(filename)[1].lower()


def is_allowed_file(filename: str) -> bool:
    return get_extension(filename) in ALLOWED_EXTENSIONS


# --------------------------------------------------------------------------
# Individual readers
# --------------------------------------------------------------------------
def _read_pdf(path: str) -> str:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            # Many PDFs are "encrypted" with an empty password; try that first.
            if reader.decrypt("") == 0:
                raise DocumentError(
                    "This PDF is password-protected. Please upload an unlocked copy."
                )
        pages = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)
        return "\n".join(pages)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError(
            "The PDF could not be read. The file may be corrupted or invalid."
        )


def _read_docx(path: str) -> str:
    try:
        document = Document(path)
    except Exception:
        raise DocumentError(
            "The DOCX file could not be read. It may be corrupted or not a real "
            "Word (.docx) file."
        )

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Tables are not part of document.paragraphs, so read them separately.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _read_txt(path: str) -> str:
    # Try the most common encodings in order.
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise DocumentError("The text file uses an unsupported character encoding.")


def _clean_text(text: str) -> str:
    """Remove null characters and collapse excessive blank lines."""
    text = text.replace("\x00", "")
    lines = [line.rstrip() for line in text.splitlines()]
    cleaned, blank = [], 0
    for line in lines:
        if line.strip() == "":
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        cleaned.append(line)
    return "\n".join(cleaned).strip()


# --------------------------------------------------------------------------
# Public function
# --------------------------------------------------------------------------
def read_document(path: str) -> str:
    """Return the cleaned text of a PDF, DOCX or TXT file."""
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        raise DocumentError("The uploaded file is empty.")

    extension = get_extension(path)
    if extension == ".pdf":
        text = _read_pdf(path)
    elif extension == ".docx":
        text = _read_docx(path)
    elif extension == ".txt":
        text = _read_txt(path)
    else:
        raise DocumentError(
            "Unsupported file type. Please upload a PDF, DOCX or TXT file."
        )

    text = _clean_text(text)
    if not text:
        if extension == ".pdf":
            raise DocumentError(
                "No text could be extracted from this PDF. It may be a scanned "
                "image-only PDF (OCR is not supported)."
            )
        raise DocumentError("The document contains no readable text.")
    return text
